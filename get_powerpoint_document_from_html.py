from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.util import Inches, Pt
from io import BytesIO

def generate_powerpoint_presentation(slides_html):
    prs = Presentation()
    
    SLIDE_LAYOUT_BLANK = 6

    for slide_content in slides_html:
        slide_layout = prs.slide_layouts[SLIDE_LAYOUT_BLANK]
        slide = prs.slides.add_slide(slide_layout)
        
        soup = BeautifulSoup(slide_content, 'html.parser')
        
        top_inch = 1.0
        left_inch = 1.0
        width_inch = 8.0
        
        for element in soup.contents:
            if not element.name:
                continue

            text = element.get_text(strip=True)
            if not text:
                continue

            if element.name == 'h1' or element.name == 'h2':
                txBox = slide.shapes.add_textbox(Inches(left_inch), Inches(top_inch), Inches(width_inch), Inches(1))
                tf = txBox.text_frame
                p = tf.paragraphs[0]
                p.text = text
                p.font.size = Pt(32 if element.name == 'h1' else 24)
                p.font.bold = True
                
                top_inch += 1.2
            
            elif element.name == 'p':
                txBox = slide.shapes.add_textbox(Inches(left_inch), Inches(top_inch), Inches(width_inch), Inches(0.5))
                tf = txBox.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = text
                p.font.size = Pt(18)
                
                top_inch += 0.8
                
            elif element.name in ['ul', 'ol']:
                txBox = slide.shapes.add_textbox(Inches(left_inch), Inches(top_inch), Inches(width_inch), Inches(2))
                tf = txBox.text_frame
                tf.word_wrap = True
                
                first_item = True
                for li in element.find_all('li'):
                    li_text = li.get_text(strip=True)
                    if first_item:
                        p = tf.paragraphs[0]
                        p.text = li_text
                        first_item = False
                    else:
                        p = tf.add_paragraph()
                        p.text = li_text
                    
                    p.font.size = Pt(18)
                    p.level = 0
                    
                top_inch += 2.0
            
            if top_inch > 7.0:
                 break

    byte_stream = BytesIO()
    prs.save(byte_stream)
    
    return byte_stream.getvalue()
